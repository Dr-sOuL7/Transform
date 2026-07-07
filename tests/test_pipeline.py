import docx

from app.core.pipeline import transform_document
from app.configs.presets import resolve_enabled


def test_roundtrip_preserves_protected_content(sample_docx, tmp_path):
    out = str(tmp_path / "out.docx")
    report = transform_document(sample_docx, out, enabled=resolve_enabled("balanced"))

    result = docx.Document(out)
    full = "\n".join(p.text for p in result.paragraphs)

    # Protected content survives verbatim.
    assert "1,234" in full
    assert "2020-01-31" in full
    assert "(Smith et al., 2019)" in full
    assert "https://example.com" in full
    assert "42%" in full

    # Filler removed, repeated transitions trimmed.
    assert "It is important to note that" not in full
    assert full.count("Moreover,") == 1
    assert full.count("Additionally,") == 1

    # Report is populated and coherent.
    assert len(report.changes) > 0
    assert report.doc_type == "academic"  # has citations
    assert not [i for i in report.issues if i.level == "error"]


def test_headings_are_not_edited(make_docx, tmp_path):
    path = make_docx([
        ("Heading 1", "It is important to note that intro"),
        (None, "It is important to note that body text follows."),
    ])
    out = str(tmp_path / "o.docx")
    transform_document(out and path, out, enabled=resolve_enabled("balanced"))
    result = docx.Document(out)
    texts = [p.text for p in result.paragraphs]
    # Heading keeps its filler (structurally protected); body is cleaned.
    assert texts[0] == "It is important to note that intro"
    assert texts[1] == "Body text follows."


def test_paragraph_count_unchanged(sample_docx, tmp_path):
    out = str(tmp_path / "out.docx")
    before = len(docx.Document(sample_docx).paragraphs)
    transform_document(sample_docx, out)
    after = len(docx.Document(out).paragraphs)
    assert before == after


def test_report_json_serialises(sample_docx, tmp_path):
    out = str(tmp_path / "out.docx")
    report = transform_document(sample_docx, out)
    payload = report.to_json()
    assert '"changes"' in payload
    assert '"summary"' in payload
