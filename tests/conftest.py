"""Shared pytest fixtures: build small .docx documents on the fly."""

from __future__ import annotations

import docx
import pytest


def _make_doc(tmp_path, paragraphs, name="sample.docx"):
    document = docx.Document()
    for style, text in paragraphs:
        if style:
            document.add_paragraph(text, style=style)
        else:
            document.add_paragraph(text)
    path = tmp_path / name
    document.save(str(path))
    return str(path)


@pytest.fixture
def sample_docx(tmp_path):
    """A document exercising filler, transitions, protected content, headings."""
    paras = [
        ("Title", "A Short Report"),
        ("Heading 1", "Introduction"),
        (None, "It is important to note that the system works offline. "
               "Moreover, it preserves meaning. Moreover, it is fast."),
        (None, "The dataset contains 1,234 records collected on 2020-01-31 "
               "(Smith et al., 2019) as reported earlier [3]. "
               "Visit https://example.com for details."),
        (None, "We reached 42% accuracy. Additionally, results were stable. "
               "Additionally, the pipeline is deterministic."),
    ]
    return _make_doc(tmp_path, paras)


@pytest.fixture
def make_docx(tmp_path):
    def _factory(paragraphs, name="doc.docx"):
        return _make_doc(tmp_path, paragraphs, name)
    return _factory
