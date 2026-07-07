"""DOCX parser: read a ``.docx`` into the internal :class:`Document` model.

The parser keeps live references to python-docx objects so the exporter can
edit run text in place and save without disturbing anything else.
"""

from __future__ import annotations

import docx
from docx.text.paragraph import Paragraph as DocxParagraph

from .model import BlockType, Document, Paragraph, Run


def _classify(docx_par: DocxParagraph) -> tuple[BlockType, int | None]:
    """Infer block type and (for lists) nesting level from the paragraph style."""
    style_name = (docx_par.style.name if docx_par.style else "") or ""
    lowered = style_name.lower()

    if lowered.startswith("title"):
        return BlockType.TITLE, None
    if lowered.startswith("heading"):
        return BlockType.HEADING, None
    if lowered.startswith("caption"):
        return BlockType.CAPTION, None
    if "quote" in lowered or "intense quote" in lowered:
        return BlockType.QUOTE, None
    if "list" in lowered:
        # Word list styles are often "List Bullet 2", "List Number", etc.
        level = None
        for token in lowered.split():
            if token.isdigit():
                level = int(token)
        return BlockType.LIST, level

    # Numbered/bulleted paragraphs sometimes rely on numbering XML rather than a
    # named list style. Treat those as list items too.
    if docx_par._p.pPr is not None and docx_par._p.pPr.numPr is not None:
        return BlockType.LIST, None

    return BlockType.BODY, None


def parse_docx(path: str) -> Document:
    """Load ``path`` and return a populated :class:`Document`."""
    docx_doc = docx.Document(path)
    document = Document(source_path=path, docx=docx_doc)

    for index, docx_par in enumerate(docx_doc.paragraphs):
        block_type, list_level = _classify(docx_par)
        runs = [Run(text=r.text, obj=r, bold=r.bold, italic=r.italic) for r in docx_par.runs]

        paragraph = Paragraph(
            index=index,
            runs=runs,
            style_name=(docx_par.style.name if docx_par.style else "") or "",
            block_type=block_type,
            obj=docx_par,
            list_level=list_level,
        )
        paragraph.original_text = paragraph.plain_text()
        document.paragraphs.append(paragraph)

    return document
